"""Markdown -> DOCX with better typography."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(r"C:\Users\Curry\Desktop\项目二\netops-assistant\docs\项目自述-从第一性原理到落地.md")
DST = Path(r"C:\Users\Curry\Desktop\项目二\netops-assistant\docs\项目自述-从第一性原理到落地.docx")

doc = Document()

# 页边距
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

# 正文样式
normal = doc.styles["Normal"]
normal.font.name = "微软雅黑"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
pf = normal.paragraph_format
pf.line_spacing = 1.6
pf.space_after = Pt(6)
pf.space_before = Pt(0)


def set_cn_font(run, name="微软雅黑"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def shade_paragraph(p, color="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            sub = re.split(r"(\*\*[^*]+\*\*)", part)
            for s in sub:
                if s.startswith("**") and s.endswith("**"):
                    r = paragraph.add_run(s[2:-2])
                    r.bold = True
                    set_cn_font(r, "微软雅黑")
                else:
                    r = paragraph.add_run(s)
                    set_cn_font(r, "微软雅黑")


lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0
in_code = False
code_buf = []
table_buf = []


def flush_table():
    global table_buf
    if not table_buf:
        return
    rows = []
    for ln in table_buf:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    rows = [r for r in rows if not all(re.match(r"^:?-+:?$", c) for c in r)]
    if rows:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Light Grid Accent 1"
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                if ci < len(t.rows[ri].cells):
                    cell.text = ""
                    p = cell.paragraphs[0]
                    add_inline(p, cell_text_safe(cell))
    table_buf = []


def cell_text_safe(_cell):
    return ""


# 重写 flush_table 让单元格能渲染 inline
def flush_table_v2():
    global table_buf
    if not table_buf:
        return
    rows = []
    for ln in table_buf:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    rows = [r for r in rows if not all(re.match(r"^:?-+:?$", c) for c in r)]
    if not rows:
        table_buf = []
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            if ci < len(t.rows[ri].cells):
                c = t.rows[ri].cells[ci]
                c.text = ""
                p = c.paragraphs[0]
                add_inline(p, cell)
    table_buf = []
    doc.add_paragraph()


while i < len(lines):
    line = lines[i]

    if line.strip().startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            shade_paragraph(p, "F5F5F5")
            r = p.add_run("\n".join(code_buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    if line.strip().startswith("|"):
        table_buf.append(line)
        i += 1
        continue
    else:
        flush_table_v2()

    if line.strip() in ("---", "***", "___"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("· · ·")
        r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        i += 1
        continue

    m = re.match(r"^(#{1,6})\s+(.*)", line)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        if level == 1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(16)
            r.bold = True
            set_cn_font(r, "微软雅黑")
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(14)
        elif level == 2:
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
            set_cn_font(r, "微软雅黑")
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.font.size = Pt(11.5)
            r.bold = True
            set_cn_font(r, "微软雅黑")
            p.paragraph_format.space_before = Pt(6)
        i += 1
        continue

    if line.strip().startswith(">"):
        # 收集连续引用行，合并成一段
        quote_lines = []
        while i < len(lines) and (lines[i].strip().startswith(">") or not lines[i].strip()):
            s = lines[i].strip()
            if s.startswith(">"):
                t = s.lstrip(">").strip()
                if t:
                    quote_lines.append(t)
            i += 1
        if quote_lines:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            text = " ".join(quote_lines)
            add_inline(p, text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                r.font.size = Pt(10.5)
        continue

    m = re.match(r"^(\s*)[-*]\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, m.group(2).strip())
        p.paragraph_format.space_after = Pt(3)
        i += 1
        continue

    m = re.match(r"^(\s*)\d+\.\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_inline(p, m.group(2).strip())
        p.paragraph_format.space_after = Pt(3)
        i += 1
        continue

    if not line.strip():
        i += 1
        continue

    p = doc.add_paragraph()
    add_inline(p, line)
    i += 1

flush_table_v2()
doc.save(DST)
print(f"OK: {DST}")
